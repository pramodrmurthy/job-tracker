"""Resume upload: validate, extract text (PDF, DOCX, TXT) and store the original file."""

import io
import re
import uuid
from pathlib import Path

from . import config

ALLOWED = {".pdf": "application/pdf",
           ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
           ".txt": "text/plain"}


class ResumeError(ValueError):
    pass


def safe_filename(name: str) -> str:
    base = Path(name or "resume").name  # drop any directory part
    stem = re.sub(r"[^A-Za-z0-9._-]+", "_", Path(base).stem).strip("._") or "resume"
    return f"{stem[:80]}{Path(base).suffix.lower()}"


def extract_text(filename: str, data: bytes) -> str:
    ext = Path(filename).suffix.lower()
    if ext not in ALLOWED:
        raise ResumeError("Unsupported file type. Upload a PDF, DOCX or TXT file.")
    if len(data) > config.MAX_RESUME_BYTES:
        raise ResumeError(f"File is larger than {config.MAX_RESUME_BYTES // (1024 * 1024)} MB.")
    if not data:
        raise ResumeError("The file is empty.")
    try:
        if ext == ".pdf":
            if not data.startswith(b"%PDF"):
                raise ResumeError("This does not look like a PDF file.")
            from pypdf import PdfReader
            text = "\n".join((page.extract_text() or "") for page in PdfReader(io.BytesIO(data)).pages)
        elif ext == ".docx":
            if not data.startswith(b"PK"):
                raise ResumeError("This does not look like a DOCX file.")
            import docx
            document = docx.Document(io.BytesIO(data))
            text = "\n".join(p.text for p in document.paragraphs)
            for table in document.tables:
                text += "\n" + "\n".join(" | ".join(c.text for c in row.cells) for row in table.rows)
        else:
            text = data.decode("utf-8", errors="replace")
    except ResumeError:
        raise
    except Exception as e:  # corrupt or encrypted file
        raise ResumeError(f"Could not read this file ({type(e).__name__}). Is it encrypted or corrupted?") from e
    text = re.sub(r"(?im)^\s*page \d+ of \d+\s*$", "", text)  # page footers add noise
    text = re.sub(r"\n{3,}", "\n\n", re.sub(r"[ \t]+", " ", text)).strip()
    if len(text) < 100:
        raise ResumeError("Almost no text could be extracted (scanned image?). Upload a text-based PDF, DOCX or TXT.")
    return text


def store_file(filename: str, data: bytes) -> Path:
    config.RESUME_DIR.mkdir(parents=True, exist_ok=True)
    path = config.RESUME_DIR / f"{uuid.uuid4().hex[:8]}_{safe_filename(filename)}"
    path.write_bytes(data)
    return path
